from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app.utils import money


def export_donors_xlsx(path: str | Path, donors: list) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    file_path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Donors"

    headers = [
        "ID",
        "Name",
        "Email",
        "Phone",
        "Address",
        "Projects",
        "Pledged",
        "Paid",
        "Remaining",
        "Donations",
        "Notes",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="176B5B")
    header_font = Font(color="FFFFFF", bold=True)
    for col, _ in enumerate(headers, start=1):
        cell = sheet.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for donor in donors:
        sheet.append(
            [
                donor.id,
                donor.name,
                donor.email or "",
                donor.phone or "",
                donor.address or "",
                len(donor.project_ids),
                float(donor.total_pledged),
                float(donor.total_paid),
                float(donor.total_remaining),
                len(donor.donations),
                donor.notes or "",
            ]
        )

    widths = [8, 22, 26, 16, 28, 10, 12, 12, 12, 12, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(file_path)
    return file_path


def export_donors_pdf(path: str | Path, donors: list, title: str = "Donors report") -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    file_path = Path(path)
    document = SimpleDocTemplate(
        str(file_path),
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#12233A"),
        spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "MetaCustom",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#6B7A90"),
        spaceAfter=12,
    )

    story = [
        Paragraph(title, title_style),
        Paragraph(
            f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} · {len(donors)} donor(s)",
            meta_style,
        ),
        Spacer(1, 4),
    ]

    data = [
        [
            "ID",
            "Name",
            "Email",
            "Phone",
            "Projects",
            "Pledged",
            "Paid",
            "Remaining",
            "Donations",
        ]
    ]
    for donor in donors:
        data.append(
            [
                str(donor.id),
                donor.name,
                donor.email or "",
                donor.phone or "",
                str(len(donor.project_ids)),
                money(donor.total_pledged),
                money(donor.total_paid),
                money(donor.total_remaining),
                str(len(donor.donations)),
            ]
        )

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#176B5B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9E3F0")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F5F8FC")],
                ),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    document.build(story)
    return file_path
